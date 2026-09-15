import argparse
import json
import os
import re
from pathlib import Path

import openpyxl


DEFAULT_SOURCE = r"C:\Users\shiyi\Desktop\JT二部RCS绿标注册\亚马逊气候友好绿标资料\JT二部已成功上标(绿标)ASIN列表-20260811更新.xlsx"


def parse_args():
    parser = argparse.ArgumentParser(description="Extract dashboard source sheets from a JT workbook.")
    parser.add_argument("source", nargs="?", default=os.environ.get("JT_SOURCE_XLSX", DEFAULT_SOURCE))
    parser.add_argument("output", nargs="?", default=os.environ.get("JT_CACHE_OUTPUT", str(Path(__file__).with_name("cached-source.json"))))
    return parser.parse_args()


def sheet_rows(book, name):
    return [[cell for cell in row] for row in book[name].iter_rows(values_only=True)]


def normalize_path_arg(value):
    text = value.strip()
    drive_path = re.search(r"[A-Za-z]:[\\/].*", text)
    if drive_path:
        text = drive_path.group(0)
    return Path(text.strip('\\"')).resolve()


args = parse_args()
source = normalize_path_arg(args.source)
output = normalize_path_arg(args.output)
book = openpyxl.load_workbook(source, data_only=True, read_only=True)
sheet_names = [name for name in ["SKU列表", "被拒列表", "成功列表", "映射关系"] if name in book.sheetnames]
payload = {
    "sourceFile": source.name,
    "sheets": {name: sheet_rows(book, name) for name in sheet_names},
}
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(
    json.dumps(payload, ensure_ascii=False, default=lambda value: value.isoformat() if hasattr(value, "isoformat") else str(value)),
    encoding="utf-8",
)
print(json.dumps({"source": str(source), "output": str(output), "sheets": {name: len(rows) for name, rows in payload["sheets"].items()}}, ensure_ascii=False))
